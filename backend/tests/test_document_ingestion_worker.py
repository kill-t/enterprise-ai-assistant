import asyncio
import hashlib
from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from uuid import uuid4

import pytest
from docx import Document
from openpyxl import Workbook

from app.common.enums import DocumentIngestionJobStatus, DocumentStatus
from app.core.config import EmbeddingSettings, RagSettings, WorkerSettings
from app.infrastructure.embedding.provider_factory import EmbeddingProviderFactory
from app.infrastructure.embedding.schemas import EmbeddingProfile
from app.infrastructure.task_queue.celery_app import celery_app
from app.modules.knowledge.processing.errors import (
    DocumentParseError,
    DocumentParseErrorCode,
)
from app.modules.knowledge.processing.ingestion_schemas import (
    DocumentIngestionSource,
    ProcessedDocumentIngestion,
)
from app.modules.knowledge.processing.parser_registry import DocumentParserRegistry
from app.modules.knowledge.processing.schemas import ChunkDraft
from app.modules.knowledge.repository.document_index_repository import (
    DocumentIndexRepository,
)
from app.modules.knowledge.repository.document_ingestion_job_repository import (
    DocumentIngestionJobRepository,
)
from app.modules.knowledge.service.document_ingestion_failure_policy import (
    DocumentIngestionFailurePolicy,
    DocumentSourceIntegrityError,
    DocumentSourceUnavailableError,
)
from app.modules.knowledge.service.document_ingestion_processor import (
    DocumentIngestionProcessor,
)
from app.modules.knowledge.service.document_ingestion_worker_service import (
    DocumentIngestionWorkerService,
)
from app.modules.knowledge.source_locator import SectionSourceLocator
from app.modules.knowledge.source_locator import SheetRangeSourceLocator


class RecordingDatabase:
    """记录 Worker 的提交和回滚次数，用于验证事务边界。"""

    def __init__(self) -> None:
        """初始化事务操作计数。"""
        self.commits = 0
        self.rollbacks = 0

    async def commit(self) -> None:
        """记录一次事务提交。"""
        self.commits += 1

    async def rollback(self) -> None:
        """记录一次事务回滚。"""
        self.rollbacks += 1


class SuccessfulIngestionProcessor:
    """模拟无需真实文件和模型的成功入库处理器。"""

    def __init__(self) -> None:
        """初始化各阶段调用记录。"""
        self.prepared_jobs = []
        self.persisted = []

    async def prepare(self, job):
        """记录已领取任务并返回普通准备数据。"""
        self.prepared_jobs.append(job)
        return SimpleNamespace(document_id=job.document_id)

    async def process(self, prepared):
        """模拟数据库事务外的下载、解析、分块和向量化。"""
        return SimpleNamespace(chunks=("chunk",), vectors=((1.0,),))

    async def persist(self, prepared, processed, finished_at) -> None:
        """记录等待与任务成功状态一起提交的索引结果。"""
        self.persisted.append((prepared, processed, finished_at))


class FailingIngestionProcessor(SuccessfulIngestionProcessor):
    """在外部处理阶段抛出指定异常。"""

    def __init__(self, error: Exception) -> None:
        """保存测试希望触发的错误。"""
        super().__init__()
        self.error = error

    async def process(self, prepared):
        """抛出预设错误，确保不会进入持久化阶段。"""
        raise self.error


def test_ingestion_job_claim_uses_single_conditional_update() -> None:
    """领取必须通过带状态、次数和重试时间条件的一条 UPDATE 完成。"""
    claimed_job = build_claimed_job()
    db = SimpleNamespace(scalar=AsyncMock(return_value=claimed_job))
    repository = DocumentIngestionJobRepository(db)

    result = asyncio.run(
        repository.claim_pending(
            job_id=claimed_job.id,
            started_at=datetime.now(timezone.utc),
        )
    )
    statement = str(db.scalar.await_args.args[0])

    assert result is claimed_job
    assert statement.startswith("UPDATE document_ingestion_jobs")
    assert "document_ingestion_jobs.status" in statement
    assert (
        "document_ingestion_jobs.attempts < document_ingestion_jobs.max_attempts"
        in statement
    )
    assert "document_ingestion_jobs.next_retry_at" in statement
    assert "attempts=(document_ingestion_jobs.attempts +" in statement


def test_rag_settings_rejects_overlap_not_smaller_than_chunk() -> None:
    """分块重叠达到 Chunk 上限时无法向前推进，配置加载阶段就应拒绝。"""
    with pytest.raises(ValueError):
        RagSettings(chunk_size_tokens=80, chunk_overlap_tokens=80)


def test_worker_settings_rejects_inverted_document_retry_delays() -> None:
    """文档重试基础延迟不能大于最大退避，否则实际等待规则会失真。"""
    with pytest.raises(ValueError):
        WorkerSettings(
            document_retry_base_delay_seconds=301,
            document_retry_max_delay_seconds=300,
        )


def test_duplicate_or_completed_ingestion_job_is_skipped() -> None:
    """没有取得执行权的重复 Celery 消息不能下载文件或写 Chunk。"""
    db = RecordingDatabase()
    job_repository = SimpleNamespace(claim_pending=AsyncMock(return_value=None))
    document_repository = SimpleNamespace(mark_processing=AsyncMock())
    processor = SuccessfulIngestionProcessor()
    service = DocumentIngestionWorkerService(
        db=db,
        job_repository=job_repository,
        document_repository=document_repository,
        processor=processor,
    )

    result = asyncio.run(service.execute(uuid4()))

    assert result == "skipped"
    document_repository.mark_processing.assert_not_awaited()
    assert processor.prepared_jobs == []
    assert db.commits == 1


def test_claimed_ingestion_job_persists_chunks_and_success_atomically() -> None:
    """索引结果、文档 READY 与任务 SUCCEEDED 应等待同一次最终提交。"""
    job = build_claimed_job(attempts=2)
    db = RecordingDatabase()
    job_repository = SimpleNamespace(
        claim_pending=AsyncMock(return_value=job),
        mark_succeeded=AsyncMock(
            return_value=SimpleNamespace(status=DocumentIngestionJobStatus.SUCCEEDED)
        ),
    )
    document_repository = SimpleNamespace(
        mark_processing=AsyncMock(return_value=job.document_id),
    )
    processor = SuccessfulIngestionProcessor()
    service = DocumentIngestionWorkerService(
        db=db,
        job_repository=job_repository,
        document_repository=document_repository,
        processor=processor,
    )

    result = asyncio.run(service.execute(job.id))

    assert result == "succeeded"
    document_repository.mark_processing.assert_awaited_once()
    job_repository.mark_succeeded.assert_awaited_once()
    assert job_repository.mark_succeeded.await_args.kwargs["expected_attempts"] == 2
    assert len(processor.persisted) == 1
    assert db.commits == 3
    assert db.rollbacks == 0


def test_permanent_parse_failure_rolls_back_and_marks_document_failed() -> None:
    """不可重试的解析错误不能保留半套 Chunk，并应同步结束任务和文档。"""
    job = build_claimed_job()
    db = RecordingDatabase()
    parse_error = DocumentParseError(
        code=DocumentParseErrorCode.PDF_ENCRYPTED,
        message="暂不支持加密 PDF",
    )
    job_repository = SimpleNamespace(
        claim_pending=AsyncMock(return_value=job),
        mark_failed=AsyncMock(
            return_value=SimpleNamespace(status=DocumentIngestionJobStatus.FAILED)
        ),
        mark_retry_pending=AsyncMock(),
    )
    document_repository = SimpleNamespace(
        mark_processing=AsyncMock(return_value=job.document_id),
        mark_failed=AsyncMock(return_value=job.document_id),
        mark_retry_pending=AsyncMock(),
    )
    service = DocumentIngestionWorkerService(
        db=db,
        job_repository=job_repository,
        document_repository=document_repository,
        processor=FailingIngestionProcessor(parse_error),
    )

    result = asyncio.run(service.execute(job.id))

    assert result == "failed"
    assert db.rollbacks == 1
    job_repository.mark_retry_pending.assert_not_awaited()
    failure = job_repository.mark_failed.await_args.kwargs
    assert failure["error_code"] == "PDF_ENCRYPTED"
    assert failure["error_message"] == "暂不支持加密 PDF"
    document_repository.mark_failed.assert_awaited_once()
    assert db.commits == 3


def test_temporary_storage_failure_returns_job_and_document_to_pending() -> None:
    """MinIO 临时不可用且次数未耗尽时，应保存延迟重试而非永久失败。"""
    job = build_claimed_job(attempts=1, max_attempts=3)
    db = RecordingDatabase()
    job_repository = SimpleNamespace(
        claim_pending=AsyncMock(return_value=job),
        mark_retry_pending=AsyncMock(
            return_value=SimpleNamespace(status=DocumentIngestionJobStatus.PENDING)
        ),
        mark_failed=AsyncMock(),
    )
    document_repository = SimpleNamespace(
        mark_processing=AsyncMock(return_value=job.document_id),
        mark_retry_pending=AsyncMock(return_value=job.document_id),
        mark_failed=AsyncMock(),
    )
    service = DocumentIngestionWorkerService(
        db=db,
        job_repository=job_repository,
        document_repository=document_repository,
        processor=FailingIngestionProcessor(DocumentSourceUnavailableError()),
        failure_policy=DocumentIngestionFailurePolicy(
            base_delay_seconds=30,
            max_delay_seconds=300,
        ),
    )

    result = asyncio.run(service.execute(job.id))

    assert result == "retry_scheduled"
    retry = job_repository.mark_retry_pending.await_args.kwargs
    assert retry["error_code"] == "DOCUMENT_SOURCE_UNAVAILABLE"
    assert retry["next_retry_at"] > retry["failed_at"]
    document_repository.mark_retry_pending.assert_awaited_once()
    job_repository.mark_failed.assert_not_awaited()


def test_stale_worker_cannot_commit_chunks_for_newer_attempt() -> None:
    """attempts 已变化时，迟到 Worker 暂存的 Chunk 和文档状态必须整体回滚。"""
    job = build_claimed_job(attempts=1)
    db = RecordingDatabase()
    job_repository = SimpleNamespace(
        claim_pending=AsyncMock(return_value=job),
        mark_succeeded=AsyncMock(return_value=None),
    )
    document_repository = SimpleNamespace(
        mark_processing=AsyncMock(return_value=job.document_id),
    )
    processor = SuccessfulIngestionProcessor()
    service = DocumentIngestionWorkerService(
        db=db,
        job_repository=job_repository,
        document_repository=document_repository,
        processor=processor,
    )

    result = asyncio.run(service.execute(job.id))

    assert result == "stale"
    assert len(processor.persisted) == 1
    assert db.commits == 2
    assert db.rollbacks == 1


def test_real_processor_verifies_source_and_builds_mock_vectors() -> None:
    """真实处理器应串起下载、TXT 解析、分块和无额度 Mock Embedding。"""
    content = "项目会议安排\n\n确认议程与参会人员。".encode()
    source = DocumentIngestionSource(
        document_id=uuid4(),
        filename="meeting.txt",
        object_key="users/u/documents/d.txt",
        size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        embedding_profile=EmbeddingProfile(
            provider="mock",
            model="mock-hash-1024",
            dimension=1024,
        ),
    )
    source_repository = SimpleNamespace(
        get_processing_source=AsyncMock(return_value=source),
        stage_success=AsyncMock(return_value=source.document_id),
    )
    storage = SimpleNamespace(get_object=AsyncMock(return_value=content))
    processor = DocumentIngestionProcessor(
        document_repository=source_repository,
        object_storage=storage,
        parser_registry=DocumentParserRegistry(),
        embedding_factory=EmbeddingProviderFactory(EmbeddingSettings()),
        chunk_size_tokens=600,
        chunk_overlap_tokens=80,
    )
    job = build_claimed_job(document_id=source.document_id)

    async def scenario():
        """依次执行准备、外部处理和暂存阶段。"""
        prepared = await processor.prepare(job)
        processed = await processor.process(prepared)
        await processor.persist(
            prepared,
            processed,
            datetime.now(timezone.utc),
        )
        return processed

    processed = asyncio.run(scenario())

    storage.get_object.assert_awaited_once_with(
        source.object_key,
        max_bytes=source.size_bytes,
    )
    assert processed.parser_version == "text-v1"
    assert processed.chunker_version == "token-paragraph-v1"
    assert processed.page_count == 0
    assert len(processed.chunks) == len(processed.vectors) >= 1
    assert all(len(vector) == 1024 for vector in processed.vectors)
    source_repository.stage_success.assert_awaited_once()


def test_real_processor_ingests_docx_table_with_mock_embedding() -> None:
    """真实处理器应让 DOCX 表格经过解析、分块和 Mock 向量化完整链路。"""
    document = Document()
    document.add_heading("资助标准", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "等级"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "一等"
    table.cell(1, 1).text = "5000 元"
    output = BytesIO()
    document.save(output)
    content = output.getvalue()
    source = DocumentIngestionSource(
        document_id=uuid4(),
        filename="资助政策.docx",
        object_key="users/u/documents/d.docx",
        size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        embedding_profile=EmbeddingProfile(
            provider="mock",
            model="mock-hash-1024",
            dimension=1024,
        ),
    )
    source_repository = SimpleNamespace(
        get_processing_source=AsyncMock(return_value=source),
        stage_success=AsyncMock(return_value=source.document_id),
    )
    processor = DocumentIngestionProcessor(
        document_repository=source_repository,
        object_storage=SimpleNamespace(get_object=AsyncMock(return_value=content)),
        parser_registry=DocumentParserRegistry(),
        embedding_factory=EmbeddingProviderFactory(EmbeddingSettings()),
        chunk_size_tokens=600,
        chunk_overlap_tokens=80,
    )

    processed = asyncio.run(processor.process(source))

    assert processed.parser_version == "docx-python-docx-v1"
    assert len(processed.chunks) == len(processed.vectors) == 1
    assert "| 等级 | 金额 |" in processed.chunks[0].content
    assert processed.chunks[0].source_locator == SectionSourceLocator(
        heading_path=("资助标准",), block_kind="table", block_index=0
    )
    assert len(processed.vectors[0]) == 1024


def test_real_processor_ingests_xlsx_range_with_mock_embedding() -> None:
    """真实处理器应让 XLSX 数据范围经过解析、分块和 Mock 向量化完整链路。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "资助标准"
    sheet.append(["等级", "金额"])
    sheet.append(["一等", 5000])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    source = DocumentIngestionSource(
        document_id=uuid4(),
        filename="资助标准.xlsx",
        object_key="users/u/documents/d.xlsx",
        size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        embedding_profile=EmbeddingProfile(
            provider="mock",
            model="mock-hash-1024",
            dimension=1024,
        ),
    )
    source_repository = SimpleNamespace(
        get_processing_source=AsyncMock(return_value=source),
        stage_success=AsyncMock(return_value=source.document_id),
    )
    processor = DocumentIngestionProcessor(
        document_repository=source_repository,
        object_storage=SimpleNamespace(get_object=AsyncMock(return_value=content)),
        parser_registry=DocumentParserRegistry(),
        embedding_factory=EmbeddingProviderFactory(EmbeddingSettings()),
        chunk_size_tokens=600,
        chunk_overlap_tokens=80,
    )

    processed = asyncio.run(processor.process(source))

    assert processed.parser_version == "xlsx-openpyxl-v1"
    assert len(processed.chunks) == len(processed.vectors) == 1
    assert "| 等级 | 金额 |" in processed.chunks[0].content
    assert processed.chunks[0].source_locator == SheetRangeSourceLocator(
        sheet_name="资助标准",
        cell_range="A2:B2",
        header_range="A1:B1",
    )
    assert len(processed.vectors[0]) == 1024


def test_processor_rejects_downloaded_content_with_different_hash() -> None:
    """对象大小相同但内容被替换时，必须在解析和 Embedding 前拒绝处理。"""
    original = b"expected"
    tampered = b"tampered"
    source = build_source(original)
    source_repository = SimpleNamespace(
        get_processing_source=AsyncMock(return_value=source),
    )
    processor = DocumentIngestionProcessor(
        document_repository=source_repository,
        object_storage=SimpleNamespace(get_object=AsyncMock(return_value=tampered)),
        parser_registry=DocumentParserRegistry(),
        embedding_factory=EmbeddingProviderFactory(EmbeddingSettings()),
        chunk_size_tokens=600,
        chunk_overlap_tokens=80,
    )

    with pytest.raises(DocumentSourceIntegrityError):
        asyncio.run(processor.process(source))


def test_index_repository_stages_chunk_replacement_and_ready_state() -> None:
    """Repository 应在同一未提交事务中删除旧索引、加入新 Chunk 并更新 READY。"""
    source = build_source("项目会议".encode("utf-8"))
    chunk = ChunkDraft(
        chunk_index=0,
        content="项目会议",
        content_hash=hashlib.sha256("项目会议".encode()).hexdigest(),
        heading_path=None,
        page_start=None,
        page_end=None,
        token_count=4,
        chunker_version="token-paragraph-v1",
    )
    processed = ProcessedDocumentIngestion(
        chunks=(chunk,),
        vectors=((1.0,) + (0.0,) * 1023,),
        embedding_model="mock-hash-1024",
        parser_version="text-v1",
        chunker_version="token-paragraph-v1",
        page_count=0,
    )
    db = SimpleNamespace(
        execute=AsyncMock(),
        scalar=AsyncMock(return_value=source.document_id),
        add_all=Mock(),
    )
    repository = DocumentIndexRepository(db)

    result = asyncio.run(
        repository.stage_success(
            source=source,
            processed=processed,
            finished_at=datetime.now(timezone.utc),
        )
    )

    delete_statement = str(db.execute.await_args.args[0])
    update_statement = db.scalar.await_args.args[0]
    update_parameters = update_statement.compile().params
    added_chunks = db.add_all.call_args.args[0]
    assert result == source.document_id
    assert delete_statement.startswith("DELETE FROM document_chunks")
    assert update_parameters["status"] == DocumentStatus.READY
    assert update_parameters["chunk_count"] == 1
    assert len(added_chunks) == 1
    assert added_chunks[0].content == "项目会议"
    assert len(added_chunks[0].embedding) == 1024


def test_celery_app_registers_document_ingestion_task() -> None:
    """Worker 启动时必须发现文档入库任务的稳定名称。"""
    celery_app.loader.import_default_modules()

    assert "knowledge.process_document_ingestion_job" in celery_app.tasks


def build_claimed_job(
    attempts: int = 1,
    max_attempts: int = 3,
    document_id=None,
):
    """构造已经被原子领取为 RUNNING 的任务替身。"""
    return SimpleNamespace(
        id=uuid4(),
        document_id=document_id or uuid4(),
        status=DocumentIngestionJobStatus.RUNNING,
        attempts=attempts,
        max_attempts=max_attempts,
    )


def build_source(content: bytes) -> DocumentIngestionSource:
    """构造大小、哈希和 Mock 模型快照一致的文档来源。"""
    return DocumentIngestionSource(
        document_id=uuid4(),
        filename="meeting.txt",
        object_key="users/u/documents/d.txt",
        size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        embedding_profile=EmbeddingProfile(
            provider="mock",
            model="mock-hash-1024",
            dimension=1024,
        ),
    )
