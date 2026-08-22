import asyncio
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

from app.modules.knowledge.processing.chunker import TokenAwareChunker
from app.modules.knowledge.processing.docx_parser import DocxParser
from app.modules.knowledge.processing.errors import (
    DocumentParseError,
    DocumentParseErrorCode,
)
from app.modules.knowledge.source_locator import SectionSourceLocator


class CharacterTokenCounter:
    """使用字符数模拟 Token，避免测试依赖具体模型 tokenizer。"""

    def count(self, text: str) -> int:
        """返回文本字符数。"""
        return len(text)


def test_docx_parser_preserves_body_order_headings_lists_and_tables() -> None:
    """DOCX 正文、表格和列表必须按原文顺序输出，并携带稳定结构位置。"""
    content = build_docx(
        lambda document: (
            document.add_paragraph("无标题前言"),
            document.add_heading("资助政策", level=1),
            document.add_paragraph("申请人须为全日制在校生。"),
            add_policy_table(document),
            document.add_heading("申请流程", level=2),
            document.add_paragraph("准备申请材料", style="List Bullet"),
        )
    )

    parsed = asyncio.run(DocxParser().parse(BytesIO(content)))

    assert parsed.parser_version == "docx-python-docx-v1"
    assert parsed.page_count == 0
    assert [section.content for section in parsed.sections] == [
        "无标题前言",
        "申请人须为全日制在校生。",
        "| 项目 | 要求 |\n| --- | --- |\n| 国籍 | 中华人民共和国 |",
        "- 准备申请材料",
    ]
    assert [section.heading_path for section in parsed.sections] == [
        (),
        ("资助政策",),
        ("资助政策",),
        ("资助政策", "申请流程"),
    ]
    assert [section.source_locator for section in parsed.sections] == [
        SectionSourceLocator(heading_path=(), block_kind="paragraph", block_index=0),
        SectionSourceLocator(
            heading_path=("资助政策",),
            block_kind="paragraph",
            block_index=1,
        ),
        SectionSourceLocator(
            heading_path=("资助政策",), block_kind="table", block_index=0
        ),
        SectionSourceLocator(
            heading_path=("资助政策", "申请流程"),
            block_kind="list",
            block_index=0,
        ),
    ]


def test_chunker_keeps_docx_block_locator_instead_of_rebuilding_it() -> None:
    """切块后必须保留 Word 表格编号，不能退化成只有标题路径的位置。"""
    content = build_docx(
        lambda document: (
            document.add_heading("资助标准", level=1),
            add_policy_table(document),
        )
    )
    parsed = asyncio.run(DocxParser().parse(BytesIO(content)))

    chunks = TokenAwareChunker(
        token_counter=CharacterTokenCounter(),
        chunk_size_tokens=200,
        chunk_overlap_tokens=20,
    ).chunk(parsed)

    assert len(chunks) == 1
    assert chunks[0].heading_path == "资助标准"
    assert chunks[0].source_locator == SectionSourceLocator(
        heading_path=("资助标准",), block_kind="table", block_index=0
    )


def test_docx_parser_groups_adjacent_paragraphs_for_token_chunking() -> None:
    """同一标题下的相邻段落应先组成来源段，避免每个短段落都成为微小 Chunk。"""
    content = build_docx(
        lambda document: (
            document.add_heading("申请条件", level=1),
            document.add_paragraph("条件一。"),
            document.add_paragraph("条件二。"),
            document.add_paragraph("条件三。"),
        )
    )

    parsed = asyncio.run(DocxParser().parse(BytesIO(content)))

    assert len(parsed.sections) == 1
    assert parsed.sections[0].content == "条件一。\n\n条件二。\n\n条件三。"
    assert parsed.sections[0].source_locator == SectionSourceLocator(
        heading_path=("申请条件",), block_kind="paragraph", block_index=0
    )


def test_docx_parser_reports_stable_invalid_error() -> None:
    """损坏的 Word 必须产生可持久化的稳定错误码。"""
    error = run_parse_error(DocxParser().parse(BytesIO(b"not-a-zip")))

    assert error.code == DocumentParseErrorCode.DOCX_INVALID


def test_docx_parser_reports_stable_empty_error() -> None:
    """没有正文的 Word 必须明确提示当前无法提取文字。"""
    content = build_docx(lambda document: None)

    error = run_parse_error(DocxParser().parse(BytesIO(content)))

    assert error.code == DocumentParseErrorCode.NO_EXTRACTABLE_TEXT


def test_docx_parser_rejects_macro_payload_even_when_extension_is_docx() -> None:
    """伪装成 docx 的宏载荷必须在交给第三方解析器前被拒绝。"""
    content = BytesIO(build_docx(lambda document: document.add_paragraph("正文")))
    with ZipFile(content, mode="a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/vbaProject.bin", b"macro")

    error = run_parse_error(DocxParser().parse(BytesIO(content.getvalue())))

    assert error.code == DocumentParseErrorCode.OFFICE_ARCHIVE_UNSAFE


def test_docx_parser_rejects_archive_over_entry_limit() -> None:
    """DOCX ZIP 条目过多时必须提前失败，避免解析不可信压缩包耗尽资源。"""
    content = build_docx(lambda document: document.add_paragraph("正文"))

    error = run_parse_error(DocxParser(max_archive_entries=1).parse(BytesIO(content)))

    assert error.code == DocumentParseErrorCode.OFFICE_ARCHIVE_UNSAFE


def add_policy_table(document) -> None:
    """向测试文档添加固定两行表格。"""
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "国籍"
    table.cell(1, 1).text = "中华人民共和国"


def build_docx(builder) -> bytes:
    """在内存中生成真实 OOXML 文件，避免测试依赖仓库二进制样本。"""
    document = Document()
    builder(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def run_parse_error(awaitable) -> DocumentParseError:
    """运行异步解析并返回预期错误。"""

    async def scenario() -> DocumentParseError:
        """捕获解析错误；未抛错时让测试明确失败。"""
        try:
            await awaitable
        except DocumentParseError as exc:
            return exc
        raise AssertionError("解析必须失败")

    return asyncio.run(scenario())
