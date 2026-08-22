"""把 XLSX 可见工作表解析成带单元格范围的 Markdown 数据块。"""

import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from itertools import zip_longest
from typing import BinaryIO
from zipfile import BadZipFile, LargeZipFile

from lxml.etree import XMLSyntaxError
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from app.modules.knowledge.processing.base import BaseDocumentParser
from app.modules.knowledge.processing.errors import (
    DocumentParseError,
    DocumentParseErrorCode,
)
from app.modules.knowledge.processing.office_archive import (
    OfficeArchiveLimits,
    OfficeArchiveValidationError,
    validate_office_archive,
)
from app.modules.knowledge.processing.schemas import ParsedDocument, ParsedSection
from app.modules.knowledge.source_locator import SheetRangeSourceLocator


REQUIRED_PACKAGE_PARTS = frozenset({"[Content_Types].xml", "xl/workbook.xml"})
EXTERNAL_LINK_PREFIXES = ("xl/externallinks/",)


@dataclass(frozen=True, slots=True)
class _SheetRow:
    """保存一个非空工作表行及其按列编号索引的规范化值。"""

    row_number: int
    values: dict[int, str]


class XlsxParser(BaseDocumentParser):
    """以只读模式提取可见工作表，并按表头与完整数据行形成来源块。"""

    parser_version = "xlsx-openpyxl-v1"

    def __init__(
        self,
        *,
        max_rows_per_block: int = 20,
        max_estimated_tokens_per_block: int = 520,
        max_worksheets: int = 50,
        max_rows_per_sheet: int = 50_000,
        max_columns_per_sheet: int = 256,
        max_non_empty_cells: int = 200_000,
        max_archive_entries: int = 2_000,
        max_uncompressed_bytes: int = 100 * 1024 * 1024,
        max_entry_uncompressed_bytes: int = 50 * 1024 * 1024,
        max_compression_ratio: int = 500,
    ) -> None:
        """保存结构切块与资源上限，防止格式正确但规模异常的工作簿耗尽资源。"""
        structural_limits = (
            max_rows_per_block,
            max_estimated_tokens_per_block,
            max_worksheets,
            max_rows_per_sheet,
            max_columns_per_sheet,
            max_non_empty_cells,
        )
        if any(limit <= 0 for limit in structural_limits):
            raise ValueError("Excel 解析与分块上限必须大于零")
        self.max_rows_per_block = max_rows_per_block
        self.max_estimated_tokens_per_block = max_estimated_tokens_per_block
        self.max_worksheets = max_worksheets
        self.max_rows_per_sheet = max_rows_per_sheet
        self.max_columns_per_sheet = max_columns_per_sheet
        self.max_non_empty_cells = max_non_empty_cells
        self.archive_limits = OfficeArchiveLimits(
            max_entries=max_archive_entries,
            max_uncompressed_bytes=max_uncompressed_bytes,
            max_entry_uncompressed_bytes=max_entry_uncompressed_bytes,
            max_compression_ratio=max_compression_ratio,
        )

    async def parse(self, stream: BinaryIO) -> ParsedDocument:
        """安全检查后同时读取公式与缓存值，输出可追溯的数据区域。"""
        stream.seek(0)
        content = stream.read()
        self._validate_archive(content)
        formula_workbook = None
        cached_workbook = None
        try:
            # read_only 避免普通模式按文件大小数十倍展开对象；keep_links=False 不加载外链缓存。
            formula_workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=False,
                keep_links=False,
            )
            cached_workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            sections = self._parse_workbook(formula_workbook, cached_workbook)
        except DocumentParseError:
            raise
        except (
            InvalidFileException,
            BadZipFile,
            LargeZipFile,
            XMLSyntaxError,
            KeyError,
            ValueError,
            TypeError,
            IndexError,
        ) as exc:
            raise self._invalid_document_error() from exc
        finally:
            if formula_workbook is not None:
                formula_workbook.close()
            if cached_workbook is not None:
                cached_workbook.close()

        if not sections:
            raise DocumentParseError(
                code=DocumentParseErrorCode.NO_EXTRACTABLE_TEXT,
                message="Excel 文件的可见工作表中没有可提取数据",
            )
        return ParsedDocument(
            sections=tuple(sections),
            page_count=0,
            parser_version=self.parser_version,
        )

    def _validate_archive(self, content: bytes) -> None:
        """调用 OOXML 共用检查，并拒绝宏和外部工作簿链接部件。"""
        try:
            validate_office_archive(
                content,
                required_parts=REQUIRED_PACKAGE_PARTS,
                limits=self.archive_limits,
                forbidden_prefixes=EXTERNAL_LINK_PREFIXES,
            )
        except OfficeArchiveValidationError as exc:
            if exc.unsafe:
                raise DocumentParseError(
                    code=DocumentParseErrorCode.OFFICE_ARCHIVE_UNSAFE,
                    message=f"Excel 文件未通过安全检查：{exc.reason}",
                ) from exc
            raise self._invalid_document_error() from exc

    def _parse_workbook(self, formula_workbook, cached_workbook) -> list[ParsedSection]:
        """按工作簿顺序处理可见 Worksheet，跳过隐藏的内部计算表。"""
        if len(formula_workbook.worksheets) > self.max_worksheets:
            raise self._limit_error("工作表数量超过限制")

        sections: list[ParsedSection] = []
        non_empty_cell_count = 0
        for formula_sheet in formula_workbook.worksheets:
            if formula_sheet.sheet_state != "visible":
                continue
            # OOXML 允许省略工作表 dimension；这时 openpyxl 只读模式会返回 None。
            # 预检查仅在尺寸已知时快速失败，尺寸未知仍由后续逐行、逐列和单元格计数兜底。
            known_max_row = formula_sheet.max_row or 0
            known_max_column = formula_sheet.max_column or 0
            if (
                known_max_row > self.max_rows_per_sheet
                or known_max_column > self.max_columns_per_sheet
            ):
                raise self._limit_error(
                    f"工作表“{formula_sheet.title}”的行列范围超过限制"
                )

            cached_sheet = cached_workbook[formula_sheet.title]
            rows, sheet_cell_count = self._read_sheet_rows(
                formula_sheet,
                cached_sheet,
            )
            non_empty_cell_count += sheet_cell_count
            if non_empty_cell_count > self.max_non_empty_cells:
                raise self._limit_error("工作簿非空单元格数量超过限制")
            sections.extend(self._build_sheet_sections(formula_sheet.title, rows))
        return sections

    def _read_sheet_rows(
        self, formula_sheet, cached_sheet
    ) -> tuple[list[_SheetRow], int]:
        """同步遍历公式版和缓存值版工作表，并保留真实行号。"""
        rows: list[_SheetRow] = []
        non_empty_cell_count = 0
        formula_rows = formula_sheet.iter_rows()
        cached_rows = cached_sheet.iter_rows()
        for row_number, (formula_row, cached_row) in enumerate(
            zip_longest(formula_rows, cached_rows, fillvalue=()),
            start=1,
        ):
            if row_number > self.max_rows_per_sheet:
                raise self._limit_error(
                    f"工作表“{formula_sheet.title}”的实际行数超过限制"
                )
            values: dict[int, str] = {}
            width = max(len(formula_row), len(cached_row))
            for column_index in range(1, width + 1):
                if column_index > self.max_columns_per_sheet:
                    raise self._limit_error(
                        f"工作表“{formula_sheet.title}”的实际列数超过限制"
                    )
                formula_cell = (
                    formula_row[column_index - 1]
                    if column_index <= len(formula_row)
                    else None
                )
                cached_cell = (
                    cached_row[column_index - 1]
                    if column_index <= len(cached_row)
                    else None
                )
                value = self._select_cell_value(formula_cell, cached_cell)
                number_format = (
                    formula_cell.number_format if formula_cell is not None else ""
                )
                text = self._format_cell_value(value, number_format=number_format)
                if text:
                    values[column_index] = text
                    non_empty_cell_count += 1
            if values:
                rows.append(_SheetRow(row_number=row_number, values=values))
        return rows, non_empty_cell_count

    @staticmethod
    def _select_cell_value(formula_cell, cached_cell):
        """公式优先使用已有缓存结果；缓存缺失时保留公式文本且绝不执行。"""
        if formula_cell is None:
            return None
        if formula_cell.data_type == "f":
            cached_value = cached_cell.value if cached_cell is not None else None
            return cached_value if cached_value is not None else formula_cell.value
        return formula_cell.value

    @staticmethod
    def _format_cell_value(value, *, number_format: str = "") -> str:
        """按原类型生成稳定文本，避免日期、布尔值和整数语义被意外改变。"""
        if value is None:
            return ""
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, datetime):
            if not XlsxParser._number_format_has_time(number_format):
                return value.date().isoformat()
            return value.isoformat(sep=" ", timespec="seconds")
        if isinstance(value, (date, time)):
            return value.isoformat()
        if isinstance(value, float):
            if not math.isfinite(value):
                return str(value)
            return format(value, ".15g")
        text = str(value).strip()
        return re.sub(r"[\t\u00a0 ]+", " ", text)

    @staticmethod
    def _number_format_has_time(number_format: str) -> bool:
        """根据 Excel 格式码中的小时、秒或 AM/PM 判断是否应展示时间部分。"""
        normalized = re.sub(r'"[^"]*"', "", number_format).casefold()
        return "h" in normalized or "s" in normalized or "am/pm" in normalized

    def _build_sheet_sections(
        self,
        sheet_name: str,
        rows: list[_SheetRow],
    ) -> list[ParsedSection]:
        """先按空行拆区域，再按全空列拆分左右并排的独立表格。"""
        sections: list[ParsedSection] = []
        for row_region in self._split_contiguous_rows(rows):
            used_columns = sorted(
                {column for row in row_region for column in row.values}
            )
            for column_group in self._split_contiguous_numbers(used_columns):
                group_rows = [
                    _SheetRow(
                        row_number=row.row_number,
                        values={
                            column: value
                            for column, value in row.values.items()
                            if column in column_group
                        },
                    )
                    for row in row_region
                ]
                non_empty_group_rows = [row for row in group_rows if row.values]
                for table_rows in self._split_contiguous_rows(non_empty_group_rows):
                    sections.extend(
                        self._build_table_sections(
                            sheet_name,
                            column_group,
                            table_rows,
                        )
                    )
        return sections

    def _build_table_sections(
        self,
        sheet_name: str,
        columns: list[int],
        rows: list[_SheetRow],
    ) -> list[ParsedSection]:
        """识别可选标题行和表头，再按行数与保守字符预算形成数据块。"""
        if not rows:
            return []
        title_row: _SheetRow | None = None
        header_row: _SheetRow | None = None
        data_rows = rows
        if len(rows) >= 2:
            if len(rows[0].values) == 1 and len(rows[1].values) >= 2:
                title_row = rows[0]
                header_row = rows[1]
                data_rows = rows[2:]
            else:
                header_row = rows[0]
                data_rows = rows[1:]

        if not data_rows:
            return [
                self._build_header_only_section(
                    sheet_name,
                    columns,
                    rows,
                )
            ]

        sections: list[ParsedSection] = []
        current_rows: list[_SheetRow] = []
        for row in data_rows:
            candidate_rows = [*current_rows, row]
            candidate_content = self._render_table(
                sheet_name,
                columns,
                candidate_rows,
                title_row=title_row,
                header_row=header_row,
            )
            if current_rows and (
                len(candidate_rows) > self.max_rows_per_block
                or self._estimate_tokens(candidate_content)
                > self.max_estimated_tokens_per_block
            ):
                sections.append(
                    self._build_data_section(
                        sheet_name,
                        columns,
                        current_rows,
                        title_row=title_row,
                        header_row=header_row,
                    )
                )
                current_rows = [row]
            else:
                current_rows = candidate_rows
        if current_rows:
            sections.append(
                self._build_data_section(
                    sheet_name,
                    columns,
                    current_rows,
                    title_row=title_row,
                    header_row=header_row,
                )
            )
        return sections

    def _build_data_section(
        self,
        sheet_name: str,
        columns: list[int],
        data_rows: list[_SheetRow],
        *,
        title_row: _SheetRow | None,
        header_row: _SheetRow | None,
    ) -> ParsedSection:
        """构造数据块正文以及独立的数据范围和表头范围。"""
        content = self._render_table(
            sheet_name,
            columns,
            data_rows,
            title_row=title_row,
            header_row=header_row,
        )
        header_rows = [row for row in (title_row, header_row) if row is not None]
        return ParsedSection(
            content=content,
            heading_path=(sheet_name,),
            source_locator=SheetRangeSourceLocator(
                sheet_name=sheet_name,
                cell_range=self._cell_range(columns, data_rows),
                header_range=(
                    self._cell_range(columns, header_rows) if header_rows else None
                ),
            ),
        )

    def _build_header_only_section(
        self,
        sheet_name: str,
        columns: list[int],
        rows: list[_SheetRow],
    ) -> ParsedSection:
        """只有一到两行说明或表头时仍保存其知识，但不伪造数据行。"""
        content = self._render_table(
            sheet_name,
            columns,
            rows,
            title_row=None,
            header_row=None,
        )
        return ParsedSection(
            content=content,
            heading_path=(sheet_name,),
            source_locator=SheetRangeSourceLocator(
                sheet_name=sheet_name,
                cell_range=self._cell_range(columns, rows),
                header_range=None,
            ),
        )

    def _render_table(
        self,
        sheet_name: str,
        columns: list[int],
        data_rows: list[_SheetRow],
        *,
        title_row: _SheetRow | None,
        header_row: _SheetRow | None,
    ) -> str:
        """重复表头生成 Markdown，使每个独立 Chunk 都能理解列语义。"""
        parts = [f"工作表：{sheet_name}"]
        if title_row is not None:
            title = " / ".join(title_row.values.values())
            parts.append(f"标题：{title}")
        if header_row is None:
            headers = [get_column_letter(column) for column in columns]
        else:
            headers = [
                header_row.values.get(column) or f"列{get_column_letter(column)}"
                for column in columns
            ]
        table_lines = [
            self._markdown_row(headers),
            self._markdown_row(["---"] * len(columns)),
        ]
        table_lines.extend(
            self._markdown_row([row.values.get(column, "") for column in columns])
            for row in data_rows
        )
        parts.append("\n".join(table_lines))
        return "\n\n".join(parts)

    @staticmethod
    def _markdown_row(values: list[str]) -> str:
        """转义单元格中的 Markdown 竖线和换行，保持列边界稳定。"""
        escaped = [value.replace("|", r"\|").replace("\n", "<br>") for value in values]
        return f"| {' | '.join(escaped)} |"

    @staticmethod
    def _estimate_tokens(content: str) -> int:
        """按每三个 UTF-8 字节一个 Token 保守预分块，最终预算仍由正式 Chunker 校验。"""
        return math.ceil(len(content.encode("utf-8")) / 3)

    @staticmethod
    def _cell_range(columns: list[int], rows: list[_SheetRow]) -> str:
        """根据连续列组和连续行块生成标准 A1 范围。"""
        return (
            f"{get_column_letter(columns[0])}{rows[0].row_number}:"
            f"{get_column_letter(columns[-1])}{rows[-1].row_number}"
        )

    @staticmethod
    def _split_contiguous_rows(rows: list[_SheetRow]) -> list[list[_SheetRow]]:
        """把行号不连续的区域拆开，避免 Locator 跨越整行空白。"""
        groups: list[list[_SheetRow]] = []
        for row in rows:
            if not groups or row.row_number != groups[-1][-1].row_number + 1:
                groups.append([row])
            else:
                groups[-1].append(row)
        return groups

    @staticmethod
    def _split_contiguous_numbers(values: list[int]) -> list[list[int]]:
        """把非空列按连续编号分组，用整列空白分隔左右并排的表格。"""
        groups: list[list[int]] = []
        for value in values:
            if not groups or value != groups[-1][-1] + 1:
                groups.append([value])
            else:
                groups[-1].append(value)
        return groups

    @staticmethod
    def _invalid_document_error() -> DocumentParseError:
        """构造不会暴露工作簿 XML 内部信息的损坏文件错误。"""
        return DocumentParseError(
            code=DocumentParseErrorCode.XLSX_INVALID,
            message="Excel 文件损坏或不是有效的 XLSX 文件",
        )

    @staticmethod
    def _limit_error(reason: str) -> DocumentParseError:
        """构造有效但超过第一版解析资源上限的不可重试错误。"""
        return DocumentParseError(
            code=DocumentParseErrorCode.XLSX_LIMIT_EXCEEDED,
            message=f"Excel 文件超过解析限制：{reason}",
        )
